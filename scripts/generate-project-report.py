from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "reports" / "AutoVerify_Project_Report_VI.pptx"

NAVY = RGBColor(15, 23, 42)
SLATE = RGBColor(51, 65, 85)
MUTED = RGBColor(100, 116, 139)
PALE = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(37, 99, 235)
BLUE_PALE = RGBColor(239, 246, 255)
PURPLE = RGBColor(124, 58, 237)
PURPLE_PALE = RGBColor(245, 243, 255)
GREEN = RGBColor(5, 150, 105)
GREEN_PALE = RGBColor(236, 253, 245)
AMBER = RGBColor(217, 119, 6)
AMBER_PALE = RGBColor(255, 251, 235)
RED = RGBColor(220, 38, 38)
RED_PALE = RGBColor(254, 242, 242)


def set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def set_text(shape, text, size=9, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.1)
    frame.margin_right = Inches(0.1)
    frame.margin_top = Inches(0.05)
    frame.margin_bottom = Inches(0.04)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return frame


def add_box(slide, x, y, w, h, text="", fill=WHITE, color=NAVY, size=9, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shape.adjustments[0] = 0.06
    except (IndexError, KeyError):
        pass
    set_fill(shape, fill)
    set_text(shape, text, size, color, bold, align)
    return shape


def add_text(slide, x, y, w, h, text, size=9, color=NAVY, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(box, text, size, color, bold, align)
    return box


def add_bullets(slide, x, y, w, h, items, size=9, color=SLATE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.06)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(4)
        paragraph.bullet = True
    return box


def add_icon(slide, kind, x, y, size=0.33, accent=BLUE, background=BLUE_PALE):
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    set_fill(badge, background)
    shapes = {
        "overview": MSO_SHAPE.UP_DOWN_ARROW,
        "architecture": MSO_SHAPE.HEXAGON,
        "overkill": MSO_SHAPE.GEAR_6,
        "leakage": MSO_SHAPE.RIGHT_ARROW,
        "database": MSO_SHAPE.CAN,
        "security": MSO_SHAPE.HEXAGON,
        "ui": MSO_SHAPE.FLOWCHART_DOCUMENT,
        "analytics": MSO_SHAPE.UP_DOWN_ARROW,
        "storage": MSO_SHAPE.CLOUD,
        "api": MSO_SHAPE.FLOWCHART_DOCUMENT,
        "roadmap": MSO_SHAPE.RIGHT_ARROW,
    }
    symbol = slide.shapes.add_shape(
        shapes.get(kind, MSO_SHAPE.OVAL),
        Inches(x + size * 0.24),
        Inches(y + size * 0.24),
        Inches(size * 0.52),
        Inches(size * 0.52),
    )
    set_fill(symbol, accent)


def add_header(slide, section, title, subtitle, icon="overview"):
    add_icon(slide, icon, 0.58, 0.3, 0.27, BLUE, BLUE_PALE)
    add_text(slide, 0.93, 0.32, 3.5, 0.18, section.upper(), 9, BLUE, True)
    add_text(slide, 0.58, 0.58, 11.8, 0.3, title, 11, NAVY, True)
    add_text(slide, 0.58, 0.94, 12.0, 0.22, subtitle, 9, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.25), Inches(12.15), Inches(0.018))
    set_fill(line, RGBColor(226, 232, 240))


def add_footer(slide, number):
    add_text(slide, 0.58, 7.15, 7.0, 0.15, "AutoVerify System · Báo cáo dự án", 9, MUTED)
    add_text(slide, 12.05, 7.15, 0.55, 0.15, f"{number:02d}", 9, MUTED, True, PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, fill, accent, icon):
    add_box(slide, x, y, w, h, fill=fill)
    add_icon(slide, icon, x + 0.18, y + 0.17, 0.33, accent, WHITE)
    add_text(slide, x + 0.6, y + 0.19, w - 0.78, 0.18, title, 10, accent, True)
    add_text(slide, x + 0.2, y + 0.64, w - 0.4, h - 0.76, body, 9, SLATE)


def arrow(slide, x, y, color=MUTED):
    item = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(0.28), Inches(0.22))
    set_fill(item, color)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Overview + architecture
    slide = prs.slides.add_slide(blank)
    add_header(slide, "01 · Overview", "AutoVerify System", "Nền tảng xác minh Overkill và Leakage cho vận hành chất lượng sản xuất.", "overview")
    add_card(slide, 0.58, 1.62, 2.92, 1.38, "Overkill", "False alarm: station judge N / NG / N/A nhưng không có defect thật.", AMBER_PALE, AMBER, "overkill")
    add_card(slide, 3.63, 1.62, 2.92, 1.38, "Leakage", "Escape risk: judge G / OK nhưng có defect Critical/Major.", RED_PALE, RED, "leakage")
    add_card(slide, 6.68, 1.62, 2.92, 1.38, "MariaDB", "Connection theo Machine/Line, query profile và pipeline mapping.", BLUE_PALE, BLUE, "database")
    add_card(slide, 9.73, 1.62, 2.92, 1.38, "Evidence", "SeaweedFS/local storage, feature JSON, images và audit trace.", GREEN_PALE, GREEN, "storage")
    add_text(slide, 0.58, 3.48, 4.2, 0.2, "Kiến trúc ba lớp", 10, NAVY, True)
    for index, (title, body, fill, accent, icon) in enumerate([
        ("Operations UI", "React 19 + TypeScript\nTabs, filters, pagination, drilldown", BLUE_PALE, BLUE, "ui"),
        ("AutoVerify API", "ASP.NET Core\nQuery engines, scheduler, EF Core, audit", PURPLE_PALE, PURPLE, "api"),
        ("Data & compute", "MariaDB · SeaweedFS\nC++ engine tách độc lập", GREEN_PALE, GREEN, "storage"),
    ]):
        x = 0.58 + index * 4.12
        add_card(slide, x, 3.82, 3.72, 1.65, title, body, fill, accent, icon)
        if index < 2:
            arrow(slide, x + 3.77, 4.53)
    add_box(slide, 0.58, 5.82, 12.05, 0.62, "Frontend chỉ phụ thuộc HTTP contracts; Backend làm chủ orchestration, MariaDB metadata, storage và audit.", PALE, NAVY, 10, True, PP_ALIGN.CENTER)
    add_footer(slide, 1)

    # 2. Overkill + analytics
    slide = prs.slides.add_slide(blank)
    add_header(slide, "02 · Overkill", "Pipeline tự động và phân tích Overkill", "Ưu tiên giảm false alarm, re-inspection và scrap không cần thiết.", "overkill")
    flow = [
        ("Schedule / Manual", "2–3h hoặc trigger", BLUE_PALE, BLUE, "roadmap"),
        ("Product query", "Judge IN N, NG, N/A", PURPLE_PALE, PURPLE, "database"),
        ("Defect deep scan", "AOI · AI · SPI\nChunk Fk_ID", AMBER_PALE, AMBER, "analytics"),
        ("Persist & verify", "Overkill hoặc\nConsistentNG", GREEN_PALE, GREEN, "storage"),
    ]
    for index, (title, body, fill, accent, icon) in enumerate(flow):
        x = 0.58 + index * 3.1
        add_card(slide, x, 1.65, 2.68, 1.38, title, body, fill, accent, icon)
        if index < 3:
            arrow(slide, x + 2.75, 2.27)
    add_box(slide, 0.58, 3.42, 4.1, 2.56, fill=PALE)
    add_icon(slide, "analytics", 0.8, 3.65, 0.36, BLUE, WHITE)
    add_text(slide, 1.28, 3.67, 2.8, 0.18, "Overkill Analytics", 10, NAVY, True)
    add_bullets(slide, 0.78, 4.1, 3.45, 1.35, ["Filter Machine/Line", "Min NG share: 0/30/50/70%", "Tách hệ AOI / AI / SPI", "Preview candidate và Key ID"], 9)
    add_box(slide, 4.98, 3.42, 7.65, 2.56, fill=WHITE)
    add_text(slide, 5.2, 3.67, 4.2, 0.18, "Ví dụ: NG share theo defect system", 10, NAVY, True)
    for i, (name, ratio, color) in enumerate([("AOI", 0.76, AMBER), ("AI", 0.58, PURPLE), ("SPI", 0.41, BLUE)]):
        y = 4.18 + i * 0.48
        add_text(slide, 5.22, y, 0.55, 0.16, name, 9, SLATE, True)
        track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.85), Inches(y), Inches(4.62), Inches(0.19))
        set_fill(track, PALE)
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.85), Inches(y), Inches(4.62 * ratio), Inches(0.19))
        set_fill(bar, color)
        add_text(slide, 10.62, y - 0.02, 0.62, 0.16, f"{ratio:.0%}", 9, color, True, PP_ALIGN.RIGHT)
    add_text(slide, 5.22, 5.55, 6.95, 0.2, "Guidance: AOI (lighting/contrast), AI (probability/class weights), SPI (paste volume/recipe).", 9, MUTED)
    add_footer(slide, 2)

    # 3. Leakage + queue/data
    slide = prs.slides.add_slide(blank)
    add_header(slide, "03 · Leakage", "Manual queue cho Leakage và mô hình dữ liệu", "Quy trình on-demand để kiểm tra sản phẩm nghi ngờ escape risk từ line, customer return hoặc audit.", "leakage")
    flow = [
        ("Paste / CSV", "ID + optional Line\n+ Mapping", BLUE_PALE, BLUE, "ui"),
        ("Batch queue", "Persist Pending\nBatch & Queue Items", PURPLE_PALE, PURPLE, "roadmap"),
        ("Resolve MariaDB", "Theo Line hoặc\nfull-scan", AMBER_PALE, AMBER, "database"),
        ("Leakage verify", "G/OK + Critical/Major\n= Escape Risk", RED_PALE, RED, "analytics"),
    ]
    for index, (title, body, fill, accent, icon) in enumerate(flow):
        x = 0.58 + index * 3.1
        add_card(slide, x, 1.62, 2.68, 1.35, title, body, fill, accent, icon)
        if index < 3:
            arrow(slide, x + 2.75, 2.24)
    models = [
        ("Connection", "Host · DB · encrypted secret\nMachine · Line", BLUE_PALE, BLUE, "database"),
        ("Pipeline Config", "Schema · judge filter\ndefect mappings · schedule", PURPLE_PALE, PURPLE, "api"),
        ("Product / Defect", "KeyId · FkId · severity\nfeatures · images", AMBER_PALE, AMBER, "analytics"),
        ("History / Audit", "Batch · status · analysis\nresult trace", GREEN_PALE, GREEN, "storage"),
    ]
    for index, (title, body, fill, accent, icon) in enumerate(models):
        x = 0.58 + (index % 2) * 6.08
        y = 3.62 + (index // 2) * 1.22
        add_card(slide, x, y, 5.65, 0.98, title, body, fill, accent, icon)
    add_footer(slide, 3)

    # 4. Management UI + Safe query
    slide = prs.slides.add_slide(blank)
    add_header(slide, "04 · Operations", "Giao diện vận hành và quản trị query an toàn", "Nhóm navigation theo nghiệp vụ, tách nhiệm vụ vận hành/chất lượng/cấu hình để giảm thao tác phân mảnh.", "ui")
    groups = [
        ("Dashboard", "KPI · Trends · Health", BLUE_PALE, BLUE, "overview"),
        ("Operations", "Overkill Ops · Leakage Ops\nExecutions · History", PURPLE_PALE, PURPLE, "roadmap"),
        ("Products & Analysis", "Products · Defects\nOverkill Analytics", AMBER_PALE, AMBER, "analytics"),
        ("Configuration & Data", "Connections · Queries\nStorage · Audit", GREEN_PALE, GREEN, "storage"),
    ]
    for index, (title, body, fill, accent, icon) in enumerate(groups):
        x = 0.58 + index * 3.08
        add_card(slide, x, 1.65, 2.68, 1.45, title, body, fill, accent, icon)
    add_box(slide, 0.58, 3.48, 5.9, 2.23, fill=BLUE_PALE)
    add_icon(slide, "database", 0.8, 3.72, 0.36, BLUE, WHITE)
    add_text(slide, 1.28, 3.74, 3.8, 0.18, "Connection governance", 10, BLUE, True)
    add_bullets(slide, 0.78, 4.13, 5.25, 1.12, ["Machine/Line tự gán theo connection", "Password encrypt bằng Data Protection / DPAPI", "Health/latency test và query profile bound vào MariaDB"], 9)
    add_box(slide, 6.73, 3.48, 5.9, 2.23, fill=PURPLE_PALE)
    add_icon(slide, "security", 6.95, 3.72, 0.36, PURPLE, WHITE)
    add_text(slide, 7.43, 3.74, 3.7, 0.18, "Safe SQL execution", 10, PURPLE, True)
    add_bullets(slide, 6.93, 4.13, 5.25, 1.12, ["Chỉ SELECT/WITH; block DML/DDL destructive", "{{parameter}} render thành parameterized query", "Giới hạn timeout, max concurrency và test query preview"], 9)
    add_box(slide, 0.58, 6.05, 12.05, 0.5, "Data grids có pagination; modal drilldown hỗ trợ xem defect JSON và evidence mà không làm overload giao diện.", PALE, NAVY, 10, True, PP_ALIGN.CENTER)
    add_footer(slide, 4)

    # 5. Evidence, API and deployment
    slide = prs.slides.add_slide(blank)
    add_header(slide, "05 · Governance", "Evidence, API surface và triển khai", "Mọi hành động vận hành có dấu vết cấu hình, kết quả và bằng chứng để phục vụ investigation.", "storage")
    chain = [
        ("Input", "Paste / CSV / Scheduler", BLUE_PALE, BLUE, "ui"),
        ("Query", "Connection + pipeline config", PURPLE_PALE, PURPLE, "database"),
        ("Evidence", "Feature JSON / images\nobject storage path", AMBER_PALE, AMBER, "storage"),
        ("Outcome", "History + audit + status", GREEN_PALE, GREEN, "roadmap"),
    ]
    for index, (title, body, fill, accent, icon) in enumerate(chain):
        x = 0.58 + index * 3.1
        add_card(slide, x, 1.63, 2.68, 1.25, title, body, fill, accent, icon)
        if index < 3:
            arrow(slide, x + 2.75, 2.17)
    endpoints = [
        ("Overkill", "configs · run · test mapping\nexecutions", BLUE_PALE, BLUE, "overkill"),
        ("Leakage", "paste · upload CSV · process\npaged queue items", RED_PALE, RED, "leakage"),
        ("Ops & Data", "import preview · scan execution\nhistory · storage", GREEN_PALE, GREEN, "api"),
    ]
    for index, (title, body, fill, accent, icon) in enumerate(endpoints):
        x = 0.58 + index * 4.1
        add_card(slide, x, 3.38, 3.65, 1.42, title, body, fill, accent, icon)
    add_box(slide, 0.58, 5.24, 12.05, 0.88, "SeaweedFS lưu JSON/ảnh defect; khi chưa cấu hình, hệ thống fallback local filesystem. API có thể embed frontend vào wwwroot để chạy 1 package.", NAVY, WHITE, 10, False, PP_ALIGN.CENTER)
    add_footer(slide, 5)

    # 6. Delivery roadmap
    slide = prs.slides.add_slide(blank)
    add_header(slide, "06 · Roadmap", "Trạng thái triển khai và bước tiếp theo", "MVP đã có nền tảng pipeline/UI; giai đoạn tiếp theo ưu tiên data production, quality tuning và hardening.", "roadmap")
    roadmap = [
        ("Implemented", ["React operations UI", "MariaDB connections + safe SQL", "Overkill scheduler/manual run", "Leakage batch/queue", "History, audit, evidence storage"], GREEN_PALE, GREEN, "overview"),
        ("Pilot next", ["Bind MariaDB production schema", "Validate AOI/AI/SPI mappings", "Run one-line Overkill pilot", "Tune anomaly thresholds", "Review leakage queue results"], BLUE_PALE, BLUE, "analytics"),
        ("Production", ["Secrets vault / RBAC", "Metrics + alerting", "Retention policy", "Load/performance testing", "Release/rollback procedure"], AMBER_PALE, AMBER, "security"),
    ]
    for index, (title, items, fill, accent, icon) in enumerate(roadmap):
        x = 0.58 + index * 4.08
        add_box(slide, x, 1.7, 3.72, 4.2, fill=fill)
        add_icon(slide, icon, x + 0.22, 1.94, 0.36, accent, WHITE)
        add_text(slide, x + 0.7, 1.97, 2.5, 0.18, f"{index + 1:02d}  {title}", 10, accent, True)
        add_bullets(slide, x + 0.24, 2.5, 3.1, 2.78, items, 9)
    add_box(slide, 0.58, 6.25, 12.05, 0.48, "Đề xuất gần nhất: cấu hình 1 production connection + test mapping AOI/AI/SPI + chạy pilot Overkill trên 1 line.", PALE, NAVY, 10, True, PP_ALIGN.CENTER)
    add_footer(slide, 6)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
